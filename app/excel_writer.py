"""Genera el archivo Excel con el formato de la hoja 'modelo finiquito'."""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADERS = [
    "Fact.",          # A
    "Fecha Fac",      # B
    "Fec.Recep.",     # C
    "C.I. Pac.",      # D
    "Nombre",         # E (paciente)
    "C.I. Tit.",      # F
    "Nombre",         # G (titular)
    "Total Fac.",     # H
    "Monto a pagar",  # I
    "Pagado",         # J
    "Referencia",     # K
    "Fecha pago",     # L
]

MONEY_FMT = "#,##0.00"
DATE_FMT = "DD/MM/YYYY"

COL_WIDTHS = [10, 12, 12, 12, 34, 12, 34, 15, 15, 15, 12, 12]


def _as_int(value: str):
    """Convierte a int si son solo dígitos; si no (p.ej. 'SC'), deja el texto."""
    v = (value or "").strip()
    if not v:
        return None
    return int(v) if v.isdigit() else v


def build_xlsx(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "modelo finiquito"

    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="D9E1F2")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")

    # Encabezado
    for idx, title in enumerate(HEADERS, start=1):
        c = ws.cell(row=1, column=idx, value=title)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center
        c.border = border

    # Filas de datos
    for r, row in enumerate(rows, start=2):
        values = [
            _as_int(row.get("fact")),         # A  Fact.
            row.get("fecha_fac"),             # B  Fecha Fac
            row.get("fec_recep"),             # C  Fec.Recep.
            _as_int(row.get("ci_pac")),       # D  C.I. Pac.
            row.get("nombre_pac"),            # E  Nombre (paciente)
            _as_int(row.get("ci_tit")),       # F  C.I. Tit.
            row.get("nombre_tit"),            # G  Nombre (titular)
            row.get("total_fac"),             # H  Total Fac.
            row.get("monto_pagar"),           # I  Monto a pagar
            row.get("pagado"),                # J  Pagado
            _as_int(row.get("referencia")),   # K  Referencia
            row.get("fecha_pago"),            # L  Fecha pago
        ]
        for c_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c_idx, value=value)
            cell.border = border
            if c_idx in (2, 12):          # fechas B y L
                cell.number_format = DATE_FMT
            elif c_idx in (8, 9, 10):     # montos H, I, J
                cell.number_format = MONEY_FMT

    for idx, width in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    ws.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
